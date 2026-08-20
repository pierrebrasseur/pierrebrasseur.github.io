# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from resp_r1 import R1, R3

SUMMARY = [
 ("Summary of the revision", None),
 (None, "This revision responds to a single diagnosis shared by Reviewer 1's comments: the article was organized around two objects at once — a regional-institutional regime and a placeless “feminist social work” said to need changing — and neither claim was therefore able to hold the paper. We have reorganized around one object."),
 (None, "The article now addresses a configuration of mandated care, specified in the Belgian and French medico-social regime. Four conditions define it: access to a person's body or life is institutionally mandated rather than chosen; that access is performed under a vocabulary of care, help or treatment; the authority to interpret what happens within it is held professionally rather than by the person who undergoes it; and the institution that would receive a complaint is also the institution that arranges the access. The contribution is a claim about how those conditions work together: forced intimacy is not one risk factor among others but the structural medium binding the normative system to the epistemic one, so that a single arrangement produces exposure to sexual violation and forecloses its nameability in the same movement. Epistemic injustice therefore operates before disclosure and not only after it — which is why interventions addressed to the reception of testimony cannot reach the mechanism."),
 (None, "Feminist social work is no longer positioned as the problem. It is positioned institutionally: in this regime the professional who qualifies a disclosure — decides what it is, whether it is one, and what happens next — is very often a social worker or an éducateur·rice spécialisé·e. That is the point at which the configuration converts a disclosure into a non-event, and it is a position of epistemic power rather than of intimate contact. The critique is addressed there, and it is specific and cited rather than general."),
 (None, "Reviewer 3's fourth comment, on generalizability, has been incorporated as a new subsection opening the Discussion. Because the four conditions are stated abstractly, the analysis travels: to the congregate settings that remain after incomplete deinstitutionalization, to community treatment orders, and to child protection, where the mandated access is relational or informational rather than physical. The same conditions specify where the mechanism should weaken. We note that this addition resolves rather than aggravates the tension between the two reviews: the medico-social regime is where the mechanism is specified, and the four conditions are what travel."),
 (None, "Structurally: canonical exposition has been condensed throughout (Background 1,094 → 1,060 words; the epistemic section 1,374 → 1,231; the expository paragraphs in the section on structural exposure cut substantially and the duplicated intersectional-methodology passages consolidated), and the space reinvested in the article's own analysis (the three-mechanism account of forced intimacy, the occupational analysis, and the scope section, all new or substantially expanded). Repetition flagged by Reviewer 1 has been removed, uncited generalizations have been given sources or cut, the decolonial gesture and its two references have been deleted, the occupational vocabulary has been specified, and the conclusion has been rewritten under its own heading so that it introduces no new material. Body length rises from 9,177 to 9,861 words, the increase being accounted for by material the reviews asked for."),

 (None, "Finally, we checked the manuscript's factual and bibliographic claims against sources rather than relying on the previous draft, and several corrections follow from that. The prevalence figures in the Introduction conflated two different outcomes and are now reported precisely: a pooled odds ratio of 1.50 for violence of any kind against disabled adults and a prevalence of 24.3% among people with mental illnesses (Hughes et al., 2012), and an odds ratio of 2.27 for sexual victimisation specifically (Mailhot Amborski et al., 2022). Our account of validisme and capacitisme was wrong and has been rewritten from Primerano (2022), who is now cited (see Reviewer 3, comment 5). The Huronia and Winterbourne View examples were misdescribed — survivor advocacy drove the Huronia litigation rather than following it, and the UK self-advocacy networks long predate Winterbourne View — and both are now stated accurately and made to carry argument rather than illustration. The community treatment order passage is anchored in the OCTET trial (Burns et al., 2013). Two reference entries were misordered alphabetically and one had the wrong issue number; both are corrected, and every in-text citation has been checked against the reference list."),
 (None, "Page references below are to the clean revised manuscript; section titles are given where they may be more robust than pagination."),
]

def add_par(doc, text, bold=False, size=11, space=6, align=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.name = 'Times New Roman'
    p.paragraph_format.space_after = Pt(space)
    if align: p.alignment = align
    return p

def cell_text(cell, text, bold=False, size=10):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size); r.font.name = 'Times New Roman'
    p.paragraph_format.space_after = Pt(4)

def add_table(doc, rows, header_label):
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    cell_text(hdr[0], '#', True)
    cell_text(hdr[1], header_label + ' — comments/suggestions (verbatim)', True)
    cell_text(hdr[2], 'Revisions made in the manuscript, with page numbers', True)
    for num, comment, response in rows:
        row = t.add_row().cells
        cell_text(row[0], str(num))
        cell_text(row[1], comment)
        cell_text(row[2], response)
    for row in t.rows:
        row.cells[0].width = Inches(0.5)
        row.cells[1].width = Inches(2.8)
        row.cells[2].width = Inches(3.7)
    return t

doc = docx.Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.7)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(11)

add_par(doc, 'AFFILIA: Feminist Inquiry in Social Work', True, 13, 2, WD_ALIGN_PARAGRAPH.CENTER)
add_par(doc, "Author's Responses to Comments by Reviewers & the Editor", True, 12, 10, WD_ALIGN_PARAGRAPH.CENTER)
add_par(doc, 'Manuscript ID: Affilia-26-053', False, 11, 2)
add_par(doc, 'Title: Cripping Feminist Social Work: Compulsory Able-Bodiedness, Epistemic Injustice, and Sexual Violence Against Disabled People', False, 11, 12)

for head, body in SUMMARY:
    if head: add_par(doc, head, True, 12, 6)
    if body: add_par(doc, body, False, 11, 6, WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_page_break()
add_par(doc, 'Reviewer #1', True, 12, 8)
add_table(doc, R1, 'Reviewer #1')

doc.add_page_break()
add_par(doc, 'Reviewer #3', True, 12, 8)
add_table(doc, R3, 'Reviewer #3')

doc.save('Affilia-26-053_R2_Author_Response_Form.docx')
print('wrote Affilia-26-053_R2_Author_Response_Form.docx')
