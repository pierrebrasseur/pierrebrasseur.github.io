# -*- coding: utf-8 -*-
"""Build R2 clean + change-highlighted docx from R1 + an ops table."""
import copy, re, sys
import docx
from docx.shared import RGBColor
from docx.oxml.ns import qn

import ops as OPS

SRC = 'R1_clean.docx'

def parse_runs(par, text, color=None):
    """Split on *italic* markers and add runs."""
    for i, chunk in enumerate(re.split(r'\*(.+?)\*', text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        if i % 2 == 1:
            r.italic = True
        if color is not None:
            r.font.color.rgb = color

def build(outpath, highlight):
    src = docx.Document(SRC)
    dst = docx.Document(SRC)
    # wipe body paragraphs from dst, keep styles/sectPr
    body = dst.element.body
    for p in list(body.findall(qn('w:p'))):
        body.remove(p)
    sectPr = body.find(qn('w:sectPr'))

    BLUE = RGBColor(0x00, 0x00, 0xCC)

    def emit_orig(idx):
        newp = copy.deepcopy(src.paragraphs[idx]._p)
        if sectPr is not None:
            sectPr.addprevious(newp)
        else:
            body.append(newp)

    def emit_new(text, style_idx, changed=True):
        """style_idx: index of an R1 paragraph whose pPr/style to copy."""
        tmpl = src.paragraphs[style_idx]._p
        newp = copy.deepcopy(tmpl)
        # strip all runs and bookmarks etc, keep pPr
        for child in list(newp):
            if child.tag != qn('w:pPr'):
                newp.remove(child)
        if sectPr is not None:
            sectPr.addprevious(newp)
        else:
            body.append(newp)
        par = docx.text.paragraph.Paragraph(newp, dst._body)
        par.style = src.paragraphs[style_idx].style
        parse_runs(par, text, BLUE if (highlight and changed) else None)

    n = len(src.paragraphs)
    for i in range(n):
        op = OPS.OPS.get(i)
        if op is None:
            emit_orig(i)
        elif op[0] == 'DELETE':
            pass
        elif op[0] == 'REPLACE':
            for t in op[1]:
                emit_new(t, i)
        elif op[0] == 'REPLACE_STYLED':
            for t, si in op[1]:
                emit_new(t, si)
        elif op[0] == 'KEEP_THEN_ADD':
            emit_orig(i)
            for t in op[1]:
                emit_new(t, op[2] if len(op) > 2 else i)
        elif op[0] == 'ADD_BEFORE':
            for t in op[1]:
                emit_new(t, op[2] if len(op) > 2 else i)
            emit_orig(i)
        else:
            raise ValueError(op[0])
    dst.save(outpath)
    print('wrote', outpath)

if __name__ == '__main__':
    build('Affilia-26-053_R2_clean.docx', False)
    build('Affilia-26-053_R2_changes_highlighted.docx', True)
